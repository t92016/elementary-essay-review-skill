"""Build editable teacher feedback from the same evidence-linked project as PNG."""
import argparse,json
from datetime import date
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def build(project,image,out,transcript=None,evidence_dir=None,config=None,extra_pages=None):
 config=config or {}
 d=Document(); sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(1)
 for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
  s=d.styles[name];s.font.name=config.get('font_family','Noto Sans CJK TC');s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),config.get('font_family','Noto Sans CJK TC'));s.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(11.5);d.styles['Normal'].paragraph_format.line_spacing=1.15;d.styles['Normal'].paragraph_format.space_after=Pt(4)
 d.styles['Title'].font.size=Pt(24);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(13)
 for sty in ['Normal','Title','Heading 1','Heading 2','Caption']:
  snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');d.styles[sty].element.get_or_add_pPr().append(snap)
 for name in ['Heading 1','Heading 2']:
  d.styles[name].paragraph_format.space_before=Pt(10);d.styles[name].paragraph_format.space_after=Pt(6)
 d.styles['Caption'].font.size=Pt(9);d.styles['Caption'].font.bold=False
 for border in list(d.styles.element.xpath('.//w:pBdr')):border.getparent().remove(border)
 d.core_properties.title=config.get('title',Path(project['source']['file']).stem+' 作文批改參考');d.core_properties.author='Codex';d.core_properties.subject='原圖核對、佳句理由與修改示範'
 foot=sec.footer.paragraphs[0];foot.alignment=2
 r=foot.add_run('批改參考  |  ');r.font.size=Pt(9)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
 items={a['id']:a for a in project['annotations']}
 def p(t):return d.add_paragraph(t)
 def add_image(path,first=False):
  from PIL import Image
  with Image.open(path) as source:
   width=min(config.get('image_width_inches',6.5),(4.8 if first else 7.0)*source.width/source.height)
  d.add_picture(str(path),width=Inches(width))
 def item(a):
  first_paragraph=len(d.paragraphs)
  labels={'praise':'佳句','suggestion':'修改建議','uncertain':'待確認','existing':'原有註記','typo':'錯字訂正'}
  pending=a['status']=='pending' and a['kind']!='existing'
  d.add_heading(a['id']+' '+('待覆核候選' if pending else labels[a['kind']]),level=2)
  p('狀態：'+a['status']+'；覆核者：'+a.get('review',{}).get('reviewer','尚未覆核'))
  p('原文：'+a.get('context',a['quote']))
  p(('好在哪裡：' if a['kind']=='praise' else '理由：')+a.get('reason',a['note']))
  if a.get('example'):p(('修改候選（未確認）：' if pending else '修改參考：')+a['example'])
  if a.get('replacement'):p(('替換候選（未確認）：' if pending else '訂正：')+a['replacement'])
  for paragraph in d.paragraphs[first_paragraph:-1]:paragraph.paragraph_format.keep_with_next=True
 d.add_paragraph(config.get('title',Path(project['source']['file']).stem+' 批改參考'),style='Title')
 p('國小作文批改小老師 本機公開版批改參考  |  '+config.get('date',date.today().isoformat()))
 p(config.get('summary','以下批註依原文與座標核對，請連同原圖和修改理由閱讀。'))
 add_image(image,first=True)
 d.add_paragraph(config.get('image_caption','紅色波浪線：佳句。紅色括線與圈選：修改建議。待確認項不新增批註。圖片右欄編號與後頁講評相同。'),style='Caption')
 p(config.get('scope','本報告只處理已提供的原文與照片；待確認項目不能視為確定錯字。'))
 for page in extra_pages or []:
  heading=d.add_heading(page['label'],level=1);heading.paragraph_format.page_break_before=True
  add_image(page['image'])
  p(page.get('note','本頁原圖與新增批註，編號對應後頁講評。'))
 sections=config.get('sections') or [
  {'title':'佳句與修改理由','ids':[a['id'] for a in project['annotations'] if a['status']=='confirmed' and a['kind'] in ['typo','praise','suggestion']]},
  {'title':'待確認與原有註記','ids':[a['id'] for a in project['annotations'] if (a['status']=='pending' or a['kind'] in ['uncertain','existing']) and a['status']!='rejected']}]
 shown=[]
 for section in sections:
  if not section['ids']:continue
  heading=d.add_heading(section['title'],level=1);heading.paragraph_format.page_break_before=not bool(shown)
  for ident in section['ids']:
   if ident not in items or ident in shown:raise ValueError('Unknown or duplicate report ID: '+ident)
   if items[ident]['status']=='rejected':raise ValueError('Rejected item cannot appear as a recommendation')
   item(items[ident]);shown.append(ident)
  crops=section.get('crops',[])
  if crops and evidence_dir:
   index=json.loads((Path(evidence_dir)/'evidence-index.json').read_text(encoding='utf-8-sig'));proof={r['id']:r for r in index}
   d.add_heading('原照片裁切對照',level=2)
   table=d.add_table(rows=2,cols=len(crops));table.autofit=False
   for col,crop_spec in enumerate(crops):
    ident=crop_spec if isinstance(crop_spec,str) else crop_spec['id']
    file=proof[ident]['target_file'] if isinstance(crop_spec,dict) and crop_spec.get('kind')=='target' else proof[ident]['files'][0]
    table.cell(0,col).text=ident+' 原圖字跡'
    table.cell(1,col).paragraphs[0].add_run().add_picture(str(Path(evidence_dir)/file),width=Inches(min(config.get('crop_width_inches',1.15),6.1/len(crops))))
  if section.get('note'):p(section['note'])
 expected={a['id'] for a in project['annotations'] if a['status']!='rejected'}
 if set(shown)!=expected:raise ValueError('Report omitted annotation IDs: '+str(expected-set(shown)))
 if config.get('unlocated'):
  d.add_heading('尚未可靠定位',level=1)
  for u in config['unlocated']:
   p(u.get('page','')+' '+u['id']+'　'+u['quote']);p('待確認：'+u['note'])
 p('以上改寫只是參考，請依學生真實經驗調整。模糊字不能因為 OCR 或語境看似合理，就直接判成錯字。')
 if transcript:
  heading=d.add_heading('依原照片核對的文字',level=1);heading.paragraph_format.page_break_before=True
  p(config.get('transcript_note','保留原字讀法與原文順序；方括註記表示不確定、插入或刪改，不代表學生原文。'))
  for row in transcript:
   p(row['id']+'　'+row['text'])
  d.add_heading('覆核範圍',level=2)
  p(config.get('verification_note','這份文字稿經看圖核對；OCR 原始輸出另存，不將人工修正稿當成 OCR 原始結果。'))
 out=Path(out);out.parent.mkdir(parents=True,exist_ok=True);d.save(out)
 print(str(out))
def build_case(manifest_file,out,config=None):
 from essay import validate
 file=Path(manifest_file);manifest=json.loads(file.read_text(encoding='utf-8-sig'));base=file.parent
 pages=manifest.get('pages',[])
 if not pages:raise ValueError('Case manifest needs pages')
 def resolve(name):return (base/name).resolve()
 merged=None;extra=[];proof=[];ids=set()
 for i,page in enumerate(pages):
  project=json.loads(resolve(page['project']).read_text(encoding='utf-8-sig'))
  validate(project,resolve(page['source']))
  image=resolve(page['image'])
  if not image.is_file():raise ValueError('Missing page image')
  for a in project['annotations']:
   if a['id'] in ids:raise ValueError('Duplicate ID across pages')
   ids.add(a['id'])
  if merged is None:merged=dict(project,annotations=[]);first_image=image
  merged['annotations'].extend(project['annotations'])
  if i:extra.append(dict(label=page['label'],image=image,note=page.get('note','')))
  if page.get('evidence_dir'):
   directory=resolve(page['evidence_dir'])
   for row in json.loads((directory/'evidence-index.json').read_text(encoding='utf-8-sig')):
    proof.append(dict(row,files=[str(directory/n) for n in row['files']],**({'target_file':str(directory/row['target_file'])} if row.get('target_file') else {})))
 # Evidence index is a report-only aggregation; review projects stay page-specific.
 import tempfile
 with tempfile.TemporaryDirectory() as temporary:
  (Path(temporary)/'evidence-index.json').write_text(json.dumps(proof,ensure_ascii=False),encoding='utf-8')
  build(merged,first_image,out,evidence_dir=temporary,config=config,extra_pages=extra)

if __name__=='__main__':
 ap=argparse.ArgumentParser();mode=ap.add_mutually_exclusive_group(required=True);mode.add_argument('--project');mode.add_argument('--case-manifest');ap.add_argument('--image');ap.add_argument('--out',required=True);ap.add_argument('--transcript');ap.add_argument('--evidence-dir');ap.add_argument('--report-config');a=ap.parse_args()
 config=json.loads(Path(a.report_config).read_text(encoding='utf-8-sig')) if a.report_config else None
 if a.case_manifest:build_case(a.case_manifest,a.out,config)
 else:
  if not a.image:ap.error('--image is required with --project')
  build(json.loads(Path(a.project).read_text(encoding='utf-8-sig')),a.image,a.out,json.loads(Path(a.transcript).read_text(encoding='utf-8-sig')) if a.transcript else None,a.evidence_dir,config)
