"""Synthetic regression tests for the public package, not OCR accuracy tests."""
from pathlib import Path
import copy, hashlib, importlib.util, json, re, subprocess, sys, unittest
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1]
SKILL=ROOT/'skills/review-elementary-essays-local'
sys.path.insert(0,str(SKILL/'scripts'))
import case_tools as g
from essay import source_info, editor
from PIL import Image, ImageDraw, ImageFont
from docx import Document
TEST=ROOT/'test-output/pipeline'
FONT='C:/Windows/Fonts/msjh.ttc'
RESULTS=[]

class PackageTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  assert not TEST.exists(), 'Preserve prior test outputs'
  TEST.mkdir(parents=True)
  font=ImageFont.truetype(FONT,30)
  a=Image.new('RGB',(630,840),'#fffef5');d=ImageDraw.Draw(a)
  for x,text in [(490,'我很高心。'),(420,'陽光像金色絲帶。'),(350,'我覺得很好。')]:
   for i,ch in enumerate(text):d.text((x,80+i*42),ch,font=font,fill='#222222')
  d.text((80,650),'老師原註',font=font,fill='#bb0000')
  d.ellipse((75,640,225,695),outline='#bb0000',width=2)
  a.save(TEST/'vertical.png')
  b=Image.new('RGB',(640,420),'white');dd=ImageDraw.Draw(b)
  dd.text((40,50),'媽媽笑著抱住我。',font=font,fill='#222222')
  dd.text((40,160),'今天很開心。',font=font,fill='#222222')
  b.save(TEST/'horizontal.png')
  def item(ident,kind,text,box,note,**kw):
   return dict(id=ident,kind=kind,status='confirmed',source_layer='student',quote=text,
    segments=[dict(text=text,box=box,origin='observed-crop')],note=note,
    review=dict(text=True,position=True,judgment=True,reviewer='合成測試設定，非教師覆核'),**kw)
  cls.p1={'version':1,'source':source_info(TEST/'vertical.png'),'reading':'vertical-right-to-left','coverage':'full','score':None,'annotations':[
   item('T1','typo','我很高心。',[490,80,32,210],'合成錯字：開心的心意使用高興。',target_box=[490,206,32,32],target_anchor=dict(segment=0,start=3,text='心',origin='observed-crop'),replacement='興',label_xy=[535,206]),
   item('P1','praise','陽光像金色絲帶。',[420,80,32,378],'以絲帶比喻陽光，讓光線的形狀更具體。'),
   item('S1','suggestion','我覺得很好。',[350,80,32,294],'補充哪一件事情讓你覺得好。',example='我第一次完成作品，心裡很有成就感。'),
   item('K1','typo','我很高心。',[490,80,32,210],'此候選尚未核對，不可以畫成訂正。',replacement='新'),
   item('U1','uncertain','〔疑字〕',[80,500,80,30],'合成不確定字，不能確定訂正。'),
   item('E1','existing','老師原註',[80,650,150,30],'原圖既有紅筆，不能當作新增成果。')
   ],'unlocated':[dict(id='U2',quote='〔尚未定位〕',note='沒有可靠座標，不猜框。')]}
  for ident in ('K1','U1'):
   row=next(x for x in cls.p1['annotations'] if x['id']==ident)
   row['status']='pending';row['review']=dict(text=False,position=False,judgment=False,reviewer='')
  cls.p1['annotations'][-1]['source_layer']='teacher-existing'
  cls.p2={'version':1,'source':source_info(TEST/'horizontal.png'),'reading':'horizontal-left-to-right','coverage':'full','score':None,'annotations':[
   item('P2','praise','媽媽笑著抱住我。',[40,50,285,40],'笑容與擁抱描寫動作，表現親情。',mark_side='below'),
   item('S2','suggestion','今天很開心。',[40,160,230,40],'補充發生什麼事情與感受的原因。',example='今天和家人一起散步，我很開心。',mark_side='below')]}
  g.dump(TEST/'vertical.json',cls.p1);g.dump(TEST/'horizontal.json',cls.p2)
  cls.case={'pages':[dict(label='直排合成頁',source='vertical.png',project='vertical.json'),dict(label='橫排合成頁',source='horizontal.png',project='horizontal.json')],
    'report':dict(title='本機公開版合成測試講評',summary='此為定位與文件輸出的合成案例，不能代表手寫辨識或語意批改準確率。',scope='兩頁合成圖片，包含既有紅筆、未確認候選與尚未定位項。')}
  g.dump(TEST/'case.json',cls.case)
  cls.manifest=g.export_case(TEST/'case.json',TEST/'results',FONT)

 def rejects(self,change):
  p=copy.deepcopy(self.p1);change(p)
  with self.assertRaises((ValueError,KeyError,TypeError)):g.strict_project(p,TEST/'vertical.png')

 def test_01_sources_and_quote(self):
  self.assertEqual(len(g.strict_project(self.p1,TEST/'vertical.png')),7)
  self.rejects(lambda p:p['source'].update(sha256='0'*64))
  self.rejects(lambda p:p['source'].update(width=600))
  self.rejects(lambda p:p['annotations'][0].update(quote='憑空新增'))
 def test_02_bad_position(self):
  self.rejects(lambda p:p['annotations'][0]['segments'][0].update(box=[629,80,32,210]))
  self.rejects(lambda p:p['annotations'][0].update(target_box=[0,0,10,10]))
 def test_03_anchor(self):
  self.rejects(lambda p:p['annotations'][0]['target_anchor'].update(text='錯'))
  self.rejects(lambda p:p['annotations'][0].pop('target_anchor'))
 def test_04_review(self):
  self.rejects(lambda p:p['annotations'][0]['review'].update(position=False))
  self.rejects(lambda p:p['annotations'][0]['review'].update(reviewer=''))
 def test_05_layer(self):
  self.rejects(lambda p:p['annotations'][0].update(source_layer='teacher-existing'))
  self.rejects(lambda p:p['annotations'][0].pop('source_layer'))
 def test_06_uncertain_and_missing_target(self):
  self.rejects(lambda p:p['annotations'][4].update(status='confirmed',review=dict(text=True,position=True,judgment=True,reviewer='test')))
  self.rejects(lambda p:p['annotations'][0].pop('target_box'))
  self.rejects(lambda p:p['annotations'][0].pop('label_xy'))
 def test_07_unlocated(self):
  self.rejects(lambda p:p['unlocated'][0].update(id='T1'))
  self.rejects(lambda p:p['unlocated'][0].update(target_box=[10,10,20,20]))
 def test_08_pending_existing_not_rendered(self):
  manifests=[g.read(p) for p in (TEST/'results').glob('page-*/*_render-manifest.json')]
  self.assertEqual(set(manifests[0]['rendered']),{'T1','P1','S1'})
  self.assertEqual(set(manifests[0]['skipped']),{'K1','U1','E1'})
 def test_09_preserve_original_pixels(self):
  for num,name in [(1,'vertical'),(2,'horizontal')]:
   raw=Image.open(TEST/(name+'.png')).convert('RGB')
   render=Image.open(TEST/'results'/('page-'+str(num))/(name+'_原圖批註.png')).convert('RGB')
   alpha=Image.open(TEST/'results'/('page-'+str(num))/(name+'_透明批註層.png')).getchannel('A')
   self.assertEqual(raw.size,render.size)
   for original,new,a in zip(raw.getdata(),render.getdata(),alpha.getdata()):
    if a==0:self.assertEqual(original,new)
 def test_10_red_replacement(self):
  layer=Image.open(TEST/'results/page-1/vertical_透明批註層.png')
  self.assertIsNotNone(layer.crop((535,206,568,242)).getchannel('A').getbbox())
 def test_11_reports_complete(self):
  self.assertEqual(self.manifest['reports']['docx']['status'],'generated')
  self.assertEqual(self.manifest['reports']['pdf']['status'],'generated')
  text='\n'.join(p.text for p in Document(TEST/'results/講評.docx').paragraphs)
  self.assertIn('U2',text);self.assertIn('待覆核候選',text);self.assertIn('替換候選（未確認）：新',text)
  self.assertIn('本機公開版',text)
  self.assertIn('修改參考',text)
  self.assertEqual(len(Document(TEST/'results/講評.docx').inline_shapes),2)
 def test_12_output_overwrite(self):
  with self.assertRaisesRegex(ValueError,'already exists'):g.export_case(TEST/'case.json',TEST/'results',FONT)
 def test_13_duplicate_case_ids(self):
  p=copy.deepcopy(self.p2);p['annotations'][0]['id']='T1'
  g.dump(TEST/'duplicate.json',p)
  c=copy.deepcopy(self.case);c['pages'][1]['project']='duplicate.json';g.dump(TEST/'duplicate-case.json',c)
  with self.assertRaisesRegex(ValueError,'unique across'):g.export_case(TEST/'duplicate-case.json',TEST/'duplicate-result',FONT)
  self.assertFalse((TEST/'duplicate-result').exists())
 def test_14_grade(self):
  report=dict(complete_work=True,score=dict(content=4,structure=3.5,language=4,reasons=dict(content='有原文例子',structure='有原文依據',language='有原文依據')))
  self.assertIn('77.5/100',g.score_summary(report,[self.p1,self.p2]))
  with self.assertRaises(ValueError):g.score_summary(report,[dict(self.p1,coverage='partial')])
  report['score']['content']=True
  with self.assertRaises(ValueError):g.score_summary(report,[self.p1])
 def test_15_no_python_packages(self):
  proc=subprocess.run([sys.executable,'-S',str(SKILL/'scripts/case_tools.py'),'probe'],capture_output=True,text=True,encoding='utf-8')
  self.assertEqual(proc.returncode,0,proc.stderr)
  result=json.loads(proc.stdout);self.assertFalse(result['png_possible']);self.assertFalse(result['ocr_performed'])
 def test_16_font_missing(self):
  with self.assertRaises(ValueError):g.export_case(TEST/'case.json',TEST/'missing-font',str(TEST/'missing.ttf'))
  self.assertFalse((TEST/'missing-font').exists())
 def test_17_editor_embeds_and_preserves_unlocated(self):
  text=(TEST/'results/page-1/editor.html').read_text(encoding='utf-8')
  self.assertNotIn('__PAYLOAD__',text);self.assertIn('data:image/png;base64,',text);self.assertIn('"U2"',text)
 def test_18_offline_export(self):
  body='\n'.join((SKILL/'scripts'/n).read_text(encoding='utf-8') for n in ('case_tools.py','build_report_pdf.py','render_teacher.py'))
  self.assertNotIn('urllib.request',body);self.assertNotIn('WinDLL',body);self.assertNotIn('pip install',body)
 def test_19_missing_pdf_is_recorded(self):
  original=importlib.util.find_spec
  with patch.object(g.importlib.util,'find_spec',side_effect=lambda name:None if name=='reportlab' else original(name)):
   result=g.export_case(TEST/'case.json',TEST/'no-pdf',FONT)
  self.assertEqual(result['reports']['pdf']['status'],'unavailable')
  self.assertEqual(result['reports']['docx']['status'],'generated')
  self.assertFalse((TEST/'no-pdf/講評.pdf').exists())
 def test_20_vertical_edge_rejected(self):
  p=copy.deepcopy(self.p1);p['annotations']=[p['annotations'][1]]
  p['annotations'][0]['segments'][0]['box']=[599,80,31,378]
  p['unlocated']=[]
  from render_teacher import render_teacher
  with self.assertRaisesRegex(ValueError,'outside source'):
   render_teacher(TEST/'vertical.png',p,TEST/'edge',FONT)

if __name__=='__main__':
 suite=unittest.defaultTestLoader.loadTestsFromTestCase(PackageTests)
 result=unittest.TextTestRunner(verbosity=2).run(suite)
 g.dump(ROOT/'test-output/pipeline-result.json',dict(tests_run=result.testsRun,failures=len(result.failures),errors=len(result.errors),success=result.wasSuccessful(),note='Synthetic local regression only; real OCR accuracy not tested'))
 sys.exit(0 if result.wasSuccessful() else 1)
